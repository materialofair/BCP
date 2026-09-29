"""FastAPI entry point. The built UI is served from the same origin when available."""
import csv
import io
import json
from zipfile import BadZipFile
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path

from fastapi import Depends, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload
from openpyxl import Workbook, load_workbook

from .db import SessionLocal, get_db, initialize_schema
from .models import Article, AuditLog, MaterialCategory, Review, RiskEvent, ScanJob, Site, SourceStatus, Supplier, utcnow
from .seed import remove_demo_data, seed_demo
from .auth import auth_middleware

ROOT = Path(__file__).resolve().parents[2]
DIST = ROOT / "frontend" / "dist"
RANK = {"gray": 0, "green": 1, "yellow": 2, "orange": 3, "red": 4}
IMPORT_COLUMNS = {
    "供应商名称": "supplier_name", "别名": "aliases", "官网": "website",
    "供货地点": "site_name", "城市": "city", "国家/地区": "country",
    "纬度": "latitude", "经度": "longitude", "供应材料": "materials",
    "地点类型": "site_type",
}


@asynccontextmanager
async def lifespan(_app: FastAPI):
    initialize_schema()
    with SessionLocal() as db:
        seed_demo(db)
    yield


app = FastAPI(title="封装材料供应链风险平台", version="0.1.0", lifespan=lifespan)
app.middleware("http")(auth_middleware)


class SiteInput(BaseModel):
    name: str = Field(min_length=1, max_length=240)
    city: str = ""
    country: str = ""
    latitude: float | None = Field(default=None, ge=-90, le=90)
    longitude: float | None = Field(default=None, ge=-180, le=180)
    materials: str = ""
    site_type: str = Field(default="工厂", max_length=40)


class SupplierInput(BaseModel):
    name: str = Field(min_length=1, max_length=240)
    aliases: str = ""
    website: str = ""
    monitored: bool = True
    sites: list[SiteInput] = Field(default_factory=list)


class SupplierPatch(BaseModel):
    name: str | None = Field(default=None, max_length=240)
    aliases: str | None = None
    website: str | None = None
    monitored: bool | None = None


class SitePatch(BaseModel):
    name: str | None = Field(default=None, max_length=240)
    city: str | None = None
    country: str | None = None
    latitude: float | None = Field(default=None, ge=-90, le=90)
    longitude: float | None = Field(default=None, ge=-180, le=180)
    materials: str | None = None
    site_type: str | None = Field(default=None, max_length=40)


class ReviewInput(BaseModel):
    status: str
    reason: str = Field(min_length=2)
    risk_level: str | None = None


def iso(value: datetime | None):
    return value.isoformat() + "Z" if value else None


def active_events(db: Session):
    return db.scalars(select(RiskEvent).where(RiskEvent.status.in_(["open", "reviewing", "confirmed"]))).all()


def latest_scan(db: Session):
    return db.scalar(select(ScanJob).order_by(ScanJob.created_at.desc()).limit(1))


def site_output(site: Site, events: list[RiskEvent], latest_job: ScanJob | None):
    own = [e for e in events if e.site_id == site.id]
    supplier_events = [e for e in events if e.supplier_id == site.supplier_id]
    highest = max(own, key=lambda e: RANK.get(e.risk_level, 0), default=None)
    supplier_highest = max(supplier_events, key=lambda e: RANK.get(e.risk_level, 0), default=None)
    current_job = latest_job and site.supplier.last_scan_job_id == latest_job.id
    coverage = ("demo" if site.supplier.is_demo else "disabled" if not site.supplier.monitored
                else "complete" if current_job and site.supplier.last_scan_coverage_complete
                else "partial" if current_job else "not_scanned" if not site.supplier.last_scanned_at else "stale")
    coverage_summary = {"demo": "虚构演示数据", "not_scanned": "尚未完成该供应商的新闻扫描",
                        "complete": "已完成本轮配置来源扫描", "partial": "部分来源已扫描，其他来源未覆盖",
                        "disabled": "已停用监测，历史风险线索仍保留", "stale": "本轮未扫描该供应商，沿用历史风险"}[coverage]
    # A no-event green state requires a complete scan; existing risks survive failed scans.
    default_risk = "green" if coverage in ("complete", "demo") else "gray"
    risk = highest.risk_level if highest else default_risk
    supplier_risk = supplier_highest.risk_level if supplier_highest else default_risk
    supplier_prefix = ("企业级线索，尚未确认该地点受影响：" if supplier_highest and supplier_highest.site_id is None
                       else "该供应商其他地点的线索：" if supplier_highest and supplier_highest.site_id != site.id
                       else "")
    return {"id": site.id, "supplier_id": site.supplier_id, "name": site.name,
            "supplier_name": site.supplier.name, "city": site.city, "country": site.country,
            "longitude": site.longitude, "latitude": site.latitude,
            "materials": [m.strip() for m in site.materials.replace("，", ",").split(",") if m.strip()],
            "site_type": site.site_type, "risk_level": risk,
            "risk_summary": highest.summary if highest else ("虚构演示地点" if site.supplier.is_demo else "暂无地点级风险线索"),
            "risk_scope": "site" if highest else None,
            "supplier_risk_level": supplier_risk,
            "supplier_risk_summary": supplier_prefix + supplier_highest.summary if supplier_highest else "暂无企业风险线索",
            "coverage_status": coverage, "coverage_summary": coverage_summary,
            "verification": highest.verification if highest else None,
            "last_scan": iso(site.supplier.last_scanned_at), "is_demo": site.supplier.is_demo}


def event_output(event: RiskEvent):
    return {"id": event.id, "title": event.title, "summary": event.summary,
            "supplier_id": event.supplier_id, "supplier_name": event.supplier.name,
            "site_id": event.site_id, "category": event.category, "risk_level": event.risk_level,
            "auto_risk_level": event.auto_risk_level,
            "verification": event.verification, "status": event.status,
            "matching_reason": event.matching_reason, "match_reason": event.matching_reason, "source": event.article.source,
            "url": "" if event.article.source_type == "demo" else event.article.url,
            "published_at": iso(event.article.published_at), "assessed_at": iso(event.assessed_at),
            "is_demo": event.supplier.is_demo}


def supplier_output(supplier: Supplier):
    return {"id": supplier.id, "name": supplier.name, "aliases": supplier.aliases,
            "website": supplier.website, "monitored": supplier.monitored,
            "is_demo": supplier.is_demo, "site_count": len(supplier.sites)}


def record_audit(db: Session, request: Request, entity_type: str, entity_id: int,
                 before: dict, after: dict):
    if before != after:
        db.add(AuditLog(entity_type=entity_type, entity_id=entity_id, action="update",
                        actor=request.state.user, before_json=json.dumps(before, ensure_ascii=False),
                        after_json=json.dumps(after, ensure_ascii=False)))


@app.get("/api/health")
def health():
    return {"status": "ok", "time": iso(utcnow())}


@app.get("/api/dashboard")
def dashboard(db: Session = Depends(get_db)):
    suppliers = db.scalars(select(Supplier).options(joinedload(Supplier.sites))).unique().all()
    events = active_events(db)
    last = latest_scan(db)
    high_ids = {e.supplier_id for e in events if e.risk_level in ("red", "orange")}
    return {"suppliers": len(suppliers), "total_sites": sum(len(s.sites) for s in suppliers),
            "high_risk": len(high_ids), "open_events": len(events),
            "last_scan": iso(last.finished_at) if last else None,
            "status": last.status if last else "not_scanned",
            "unverified_events": sum(e.verification == "unverified" for e in events),
            "demo_mode": any(s.is_demo for s in suppliers)}


@app.get("/api/sites")
@app.get("/api/map/sites")
def sites(db: Session = Depends(get_db)):
    records = db.scalars(select(Site).options(joinedload(Site.supplier)).order_by(Site.id)).all()
    events = active_events(db)
    last = latest_scan(db)
    return [site_output(s, events, last) for s in records]


@app.get("/api/materials")
def materials(db: Session = Depends(get_db)):
    return [m.name for m in db.scalars(select(MaterialCategory).order_by(MaterialCategory.name))]


@app.get("/api/news")
@app.get("/api/news/headlines")
def news(db: Session = Depends(get_db)):
    rows = db.scalars(select(RiskEvent).options(joinedload(RiskEvent.article), joinedload(RiskEvent.supplier))
                      .where(RiskEvent.status.in_(["open", "reviewing", "confirmed"]))
                      .order_by(RiskEvent.assessed_at.desc()).limit(100)).all()
    return [event_output(e) for e in rows]


@app.get("/api/events")
def events(db: Session = Depends(get_db)):
    rows = db.scalars(select(RiskEvent).options(joinedload(RiskEvent.article), joinedload(RiskEvent.supplier))
                      .order_by(RiskEvent.assessed_at.desc()).limit(500)).all()
    return [event_output(e) for e in rows]


@app.get("/api/events/{event_id}")
def event_detail(event_id: int, db: Session = Depends(get_db)):
    row = db.get(RiskEvent, event_id)
    if not row:
        raise HTTPException(404, "Risk event not found")
    data = event_output(row)
    data["reviews"] = [{"id": r.id, "status": r.status, "reason": r.reason,
                        "previous_risk_level": r.previous_risk_level, "risk_level": r.risk_level,
                        "reviewer": r.reviewer, "created_at": iso(r.created_at)}
                       for r in db.scalars(select(Review).where(Review.event_id == event_id).order_by(Review.created_at))]
    return data


@app.post("/api/events/{event_id}/reviews")
def review(event_id: int, item: ReviewInput, request: Request, db: Session = Depends(get_db)):
    row = db.get(RiskEvent, event_id)
    if not row:
        raise HTTPException(404, "Risk event not found")
    if item.status not in {"open", "reviewing", "confirmed", "false_positive", "resolved"}:
        raise HTTPException(422, "Invalid review status")
    if item.risk_level and item.risk_level not in {"red", "orange", "yellow", "green"}:
        raise HTTPException(422, "Invalid risk level")
    if item.risk_level == "red" and item.status != "confirmed":
        raise HTTPException(422, "Red risk requires confirmed review")
    if row.auto_risk_level is None:
        row.auto_risk_level = row.risk_level
    previous = row.risk_level
    db.add(Review(event_id=event_id, status=item.status, reason=item.reason,
                  reviewer=request.state.user, previous_risk_level=previous, risk_level=item.risk_level or previous))
    row.status = item.status
    row.verification = "human_reviewed"
    if item.risk_level:
        row.risk_level = item.risk_level
    db.commit()
    return event_output(row)


@app.get("/api/suppliers")
def suppliers(db: Session = Depends(get_db)):
    rows = db.scalars(select(Supplier).options(joinedload(Supplier.sites)).order_by(Supplier.id)).unique().all()
    return [supplier_output(s) for s in rows]


@app.get("/api/suppliers/{supplier_id}")
def supplier_detail(supplier_id: int, db: Session = Depends(get_db)):
    row = db.get(Supplier, supplier_id)
    if not row:
        raise HTTPException(404, "Supplier not found")
    result = supplier_output(row)
    result["sites"] = [{"id": s.id, "name": s.name, "city": s.city, "country": s.country,
                        "latitude": s.latitude, "longitude": s.longitude, "materials": s.materials}
                       for s in row.sites]
    return result


@app.post("/api/suppliers", status_code=201)
def add_supplier(item: SupplierInput, db: Session = Depends(get_db)):
    if db.scalar(select(Supplier.id).where(Supplier.name == item.name)):
        raise HTTPException(409, "Supplier already exists")
    remove_demo_data(db)
    row = Supplier(name=item.name.strip(), aliases=item.aliases, website=item.website, monitored=item.monitored)
    db.add(row)
    db.flush()
    for s in item.sites:
        db.add(Site(supplier_id=row.id, **s.model_dump()))
    db.commit()
    db.refresh(row)
    return supplier_output(row)


@app.post("/api/suppliers/{supplier_id}/sites", status_code=201)
def add_site(supplier_id: int, item: SiteInput, db: Session = Depends(get_db)):
    if not db.get(Supplier, supplier_id):
        raise HTTPException(404, "Supplier not found")
    if db.scalar(select(Site.id).where(Site.supplier_id == supplier_id, Site.name == item.name)):
        raise HTTPException(409, "Site already exists")
    row = Site(supplier_id=supplier_id, **item.model_dump())
    db.add(row)
    db.commit()
    return {"id": row.id}


@app.patch("/api/suppliers/{supplier_id}")
def update_supplier(supplier_id: int, item: SupplierPatch, request: Request, db: Session = Depends(get_db)):
    row = db.get(Supplier, supplier_id)
    if not row:
        raise HTTPException(404, "Supplier not found")
    changes = item.model_dump(exclude_unset=True)
    if any(value is None for value in changes.values()):
        raise HTTPException(422, "Supplier fields cannot be null")
    if "name" in changes:
        changes["name"] = changes["name"].strip()
        if not changes["name"]:
            raise HTTPException(422, "Supplier name cannot be blank")
        if db.scalar(select(Supplier.id).where(Supplier.name == changes["name"], Supplier.id != supplier_id)):
            raise HTTPException(409, "Supplier already exists")
    before = {key: getattr(row, key) for key in changes}
    for key, value in changes.items():
        setattr(row, key, value)
    record_audit(db, request, "supplier", supplier_id, before, changes)
    db.commit()
    return supplier_output(row)


@app.patch("/api/sites/{site_id}")
def update_site(site_id: int, item: SitePatch, request: Request, db: Session = Depends(get_db)):
    row = db.get(Site, site_id)
    if not row:
        raise HTTPException(404, "Site not found")
    changes = item.model_dump(exclude_unset=True)
    if any(value is None for key, value in changes.items() if key not in ("latitude", "longitude")):
        raise HTTPException(422, "Site text fields cannot be null")
    if "name" in changes:
        changes["name"] = changes["name"].strip()
        if not changes["name"]:
            raise HTTPException(422, "Site name cannot be blank")
        if db.scalar(select(Site.id).where(Site.supplier_id == row.supplier_id,
                                           Site.name == changes["name"], Site.id != site_id)):
            raise HTTPException(409, "Site already exists")
    before = {key: getattr(row, key) for key in changes}
    for key, value in changes.items():
        setattr(row, key, value)
    record_audit(db, request, "site", site_id, before, changes)
    db.commit()
    last = latest_scan(db)
    return site_output(row, active_events(db), last)


@app.get("/api/supplier-import-template.xlsx")
def supplier_import_template():
    book = Workbook()
    sheet = book.active
    sheet.title = "供应商与供货地点"
    sheet.append(list(IMPORT_COLUMNS))
    sheet.freeze_panes = "A2"
    for column, width in {"A": 26, "B": 28, "C": 32, "D": 26, "E": 18,
                          "F": 18, "G": 14, "H": 14, "I": 32, "J": 18}.items():
        sheet.column_dimensions[column].width = width
    notes = book.create_sheet("填写说明")
    notes.append(["填写说明", "每行填写一个供货地点；同一供应商有多个地点时，重复填写供应商名称。"])
    notes.append(["必填", "供应商名称；有地点时填写供货地点。"])
    notes.append(["别名", "多个名称可用英文逗号、中文逗号或 | 分隔。"])
    notes.append(["坐标", "纬度 -90～90，经度 -180～180；未知时留空，不会放到国家中心。"])
    notes.append(["文件格式", "请保存为 .xlsx；不要删除第一行表头。"])
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 90
    output = io.BytesIO()
    book.save(output)
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=supplier-import-template.xlsx"})


@app.post("/api/suppliers/import")
async def import_suppliers(file: UploadFile = File(...), db: Session = Depends(get_db)):
    raw = await file.read(2_000_001)
    if len(raw) > 2_000_000:
        raise HTTPException(413, "File exceeds 2 MB")
    if (file.filename or "").lower().endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            sheet = book.active
            values = sheet.values
            headers = [IMPORT_COLUMNS.get(str(v).strip(), str(v).strip()) if v is not None else ""
                       for v in next(values)]
            reader = (dict(zip(headers, ["" if v is None else str(v) for v in row])) for row in values)
        except (BadZipFile, ValueError, StopIteration) as exc:
            raise HTTPException(422, f"Invalid XLSX: {exc}")
    elif (file.filename or "").lower().endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(422, "CSV must be UTF-8")
        csv_reader = csv.DictReader(io.StringIO(text))
        headers = [IMPORT_COLUMNS.get(h.strip(), h.strip()) for h in (csv_reader.fieldnames or [])]
        csv_reader.fieldnames = headers
        reader = csv_reader
    else:
        raise HTTPException(422, "Use an .xlsx Excel file or UTF-8 .csv file")
    if "supplier_name" not in headers:
        raise HTTPException(422, "Required header: 供应商名称 or supplier_name")
    named_headers = [header for header in headers if header]
    if len(named_headers) != len(set(named_headers)):
        raise HTTPException(422, "Duplicate column headers")
    created = updated = 0
    created_ids: set[int] = set()
    updated_ids: set[int] = set()
    problems = []
    for line, data in enumerate(reader, start=2):
        name = (data.get("supplier_name") or "").strip()
        if not name:
            problems.append(f"Line {line}: supplier_name missing")
            continue
        try:
            lat = float(data["latitude"]) if (data.get("latitude") or "").strip() else None
            lon = float(data["longitude"]) if (data.get("longitude") or "").strip() else None
            if lat is not None and not -90 <= lat <= 90 or lon is not None and not -180 <= lon <= 180:
                raise ValueError("coordinates out of range")
        except ValueError as exc:
            problems.append(f"Line {line}: {exc}")
            continue
        if len((data.get("site_type") or "").strip()) > 40:
            problems.append(f"Line {line}: site_type exceeds 40 characters")
            continue
        supplier = db.scalar(select(Supplier).where(Supplier.name == name))
        if not supplier or supplier.is_demo:
            remove_demo_data(db)
            supplier = Supplier(name=name, aliases=(data.get("aliases") or "").strip(),
                                website=(data.get("website") or "").strip())
            db.add(supplier)
            db.flush()
            created_ids.add(supplier.id)
            created += 1
        else:
            if supplier.id not in created_ids and supplier.id not in updated_ids:
                updated_ids.add(supplier.id)
                updated += 1
            if (data.get("aliases") or "").strip():
                supplier.aliases = data["aliases"].strip()
            if (data.get("website") or "").strip():
                supplier.website = data["website"].strip()
        site_name = (data.get("site_name") or "").strip()
        if site_name:
            site = db.scalar(select(Site).where(Site.supplier_id == supplier.id, Site.name == site_name))
            if not site:
                site = Site(supplier_id=supplier.id, name=site_name)
                db.add(site)
            for field in ("city", "country", "materials", "site_type"):
                value = (data.get(field) or "").strip()
                if value:
                    setattr(site, field, value)
            if lat is not None:
                site.latitude = lat
            if lon is not None:
                site.longitude = lon
    db.commit()
    return {"created": created, "updated": updated, "errors": problems}


@app.post("/api/scan-jobs", status_code=202)
def enqueue_scan(db: Session = Depends(get_db)):
    row = ScanJob(status="queued", trigger="manual")
    db.add(row)
    db.commit()
    return {"id": row.id, "status": row.status}


@app.get("/api/scan-jobs")
def scan_jobs(db: Session = Depends(get_db)):
    rows = db.scalars(select(ScanJob).order_by(ScanJob.created_at.desc()).limit(50)).all()
    return [{"id": r.id, "status": r.status, "trigger": r.trigger,
             "created_at": iso(r.created_at), "started_at": iso(r.started_at),
             "finished_at": iso(r.finished_at), "completed_at": iso(r.finished_at),
             "articles_found": r.articles_found,
             "events_created": r.events_created, "errors": r.errors, "error": r.errors} for r in rows]


@app.get("/api/source-status")
def source_status(db: Session = Depends(get_db)):
    rows = db.scalars(select(SourceStatus).order_by(SourceStatus.name)).all()
    sources = [{"name": r.name, "status": r.status, "checked_at": iso(r.checked_at),
                "last_success_at": iso(r.last_success_at), "message": r.message} for r in rows]
    return {"status": "ok" if any(s["status"] == "ok" for s in sources) else "not_scanned" if not sources else "failed",
            "message": f"{sum(s['status'] == 'ok' for s in sources)}/{len(sources)} 个来源可用" if sources else "尚未运行采集任务",
            "sources": sources}


if DIST.exists():
    app.mount("/assets", StaticFiles(directory=DIST / "assets"), name="assets")
    if (DIST / "maps").exists():
        app.mount("/maps", StaticFiles(directory=DIST / "maps"), name="maps")

    @app.get("/{path:path}", include_in_schema=False)
    def spa(path: str):
        candidate = (DIST / path).resolve()
        if candidate.is_relative_to(DIST) and candidate.is_file():
            return FileResponse(candidate)
        return FileResponse(DIST / "index.html")
